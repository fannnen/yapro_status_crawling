# main.py

import os
import shutil
import argparse
from typing import List, Optional, Callable

from openpyxl import load_workbook

from GC335 import GC335Client
from GC337 import GC337Client


PLATE_COL = "D"


def safe_save_workbook(wb, filename):
    try:
        wb.save(filename)
        print(f"\nSaved successfully: {filename}")

    except PermissionError:
        base, ext = os.path.splitext(filename)
        out_name = f"{base}_out{ext}"
        wb.save(out_name)

        print(f"\n'{filename}' is locked, probably open in Excel.")
        print(f"Saved to: {out_name}")


def create_backup_copy(file_path: str) -> str:
    output_folder = os.path.dirname(os.path.dirname(file_path))
    os.makedirs(output_folder, exist_ok=True)

    base_name = os.path.basename(file_path)
    name, ext = os.path.splitext(base_name)

    copy_path = os.path.join(output_folder, f"{name}_copy{ext}")

    counter = 1
    while os.path.exists(copy_path):
        copy_path = os.path.join(output_folder, f"{name}_copy{counter}{ext}")
        counter += 1

    shutil.copy2(file_path, copy_path)
    print(f"Backup created: {copy_path}")

    return copy_path


def read_column_values(
    excel_path: str,
    sheet_name: str,
    col_letter: str,
    start_row: int = 2
) -> List[str]:

    wb = load_workbook(excel_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"
        )

    ws = wb[sheet_name]
    col_letter = col_letter.strip().upper()

    values = []

    for cell in ws[col_letter][start_row - 1:]:
        if cell.value is None:
            continue

        s = str(cell.value).strip()

        if s:
            values.append(s)

    wb.close()
    return values


def detect_vehicle_type(ws) -> str:
    first_plate = ws[f"{PLATE_COL}2"].value

    if first_plate is None or str(first_plate).strip() == "":
        raise ValueError(f"No license plate found at {PLATE_COL}2.")

    plate = str(first_plate).strip().upper()

    if plate.startswith("E") or plate.startswith("RE"):
        return "337"

    return "335"


def find_row_by_value(ws, value, start_row=2):
    target = str(value).strip()

    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            if cell.value is None:
                continue

            if str(cell.value).strip() == target:
                return cell.row

    return None


def first_empty_row_any(ws, start_row=2):
    r = start_row
    check_cols = max(ws.max_column, 10)

    while True:
        any_value = False

        for c in range(1, check_cols + 1):
            if ws.cell(row=r, column=c).value not in (None, ""):
                any_value = True
                break

        if not any_value:
            return r

        r += 1


def first_empty_col_in_row(ws, row, start_col=1):
    for c in range(start_col, ws.max_column + 1):
        if ws.cell(row=row, column=c).value in (None, ""):
            return c

    return ws.max_column + 1


def get_default_sheet_name(excel_path: str) -> str:
    wb = load_workbook(excel_path, read_only=True)
    sheet_name = wb.sheetnames[0]
    wb.close()
    return sheet_name


def is_excel_ready(file_path: str) -> bool:
    try:
        wb = load_workbook(file_path, read_only=True)
        wb.close()
        return True

    except Exception as e:
        print(f"File not ready yet, skipping: {file_path}")
        print(f"Reason: {e}")
        return False


def should_skip_file(file_path: str) -> bool:
    filename = os.path.basename(file_path)

    if filename.startswith("~$"):
        return True

    lower = filename.lower()

    if not lower.endswith(".xlsx"):
        return True

    if "_copy" in lower:
        return True

    if "_out" in lower:
        return True

    if "_done" in lower:
        return True

    if "_updated" in lower:
        return True

    return False


def move_to_output(file_path: str, output_folder: str) -> str:
    os.makedirs(output_folder, exist_ok=True)

    base_name = os.path.basename(file_path)
    name, ext = os.path.splitext(base_name)

    dest_path = os.path.join(output_folder, f"{name}_updated{ext}")

    counter = 1
    while os.path.exists(dest_path):
        dest_path = os.path.join(output_folder, f"{name}_updated{counter}{ext}")
        counter += 1

    shutil.move(file_path, dest_path)

    print(f"Moved updated file to: {dest_path}")

    return dest_path


def run_job(
    excel_path: str,
    sheet_name: str,
    headless: bool = True,
    progress_cb: Optional[Callable[[int, int, str], None]] = None,
):

    excel_path = excel_path.strip().strip('"').strip("'")

    if not excel_path.lower().endswith(".xlsx"):
        raise ValueError("Please select a .xlsx file")

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"File not found: {excel_path}")

    create_backup_copy(excel_path)

    wb = load_workbook(excel_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"
        )

    ws = wb[sheet_name]

    plates = read_column_values(
        excel_path,
        sheet_name,
        PLATE_COL,
        start_row=2
    )

    total = len(plates)

    if total == 0:
        raise ValueError(f"No license plates found in column {PLATE_COL}.")

    if progress_cb:
        progress_cb(0, total, "")

    vehicle_type = detect_vehicle_type(ws)

    print(f"Detected vehicle type: {vehicle_type}")

    if vehicle_type == "337":
        client = GC337Client(headless=headless)

        keys_order = [
            "transId",
            "crtDate",
            "status",
            "refundDt",
        ]

    else:
        client = GC335Client(headless=headless)

        keys_order = [
            "receiveDate",
            "dealNote",
            "taxreFund",
            "custCd",
            "caseNo",
            "msg",
            "dealDate",
            "rptDate",
        ]

    try:
        for idx, plate in enumerate(plates, start=1):

            if progress_cb:
                progress_cb(idx, total, plate)

            print(f"\n[{idx}/{total}] Querying: {plate}")

            row = find_row_by_value(ws, plate, start_row=2)

            if row is None:
                row = first_empty_row_any(ws, start_row=2)
                ws.cell(row=row, column=1, value=plate)

            json_data = client.query_plate_first_row(plate)

            print(f"Result: {json_data}")

            no_result = (
                (not isinstance(json_data, dict))
                or (not json_data)
                or (json_data.get("ok") is False)
            )

            start_col = first_empty_col_in_row(
                ws,
                row,
                start_col=1
            )

            if no_result:
                ws.cell(row=row, column=start_col, value="無結果")

                reason = (
                    json_data.get("error", "")
                    if isinstance(json_data, dict)
                    else ""
                )

                ws.cell(row=row, column=start_col + 1, value=reason)

                continue

            for i, key in enumerate(keys_order):
                ws.cell(
                    row=row,
                    column=start_col + i,
                    value=json_data.get(key, "")
                )

    finally:
        try:
            client.close()
        except Exception:
            pass

    safe_save_workbook(wb, excel_path)
    wb.close()


def run_folder_mode(
    folder_path: str,
    headless: bool = True,
    move_done: bool = True,
):

    folder_path = folder_path.strip().strip('"').strip("'")

    if not os.path.isdir(folder_path):
        raise FileNotFoundError(f"Folder not found: {folder_path}")

    output_folder = os.path.dirname(folder_path)

    files = [
        os.path.join(folder_path, f)
        for f in os.listdir(folder_path)
        if not should_skip_file(os.path.join(folder_path, f))
    ]

    if not files:
        print("No Excel files found.")
        return

    # Select only the newest Excel file in the folder
    file_path = max(files, key=os.path.getmtime)

    print("\n==============================")
    print(f"Newest file selected: {file_path}")

    if not is_excel_ready(file_path):
        return

    try:
        sheet_name = get_default_sheet_name(file_path)

        print(f"Using sheet: {sheet_name}")

        run_job(
            excel_path=file_path,
            sheet_name=sheet_name,
            headless=headless,
            progress_cb=None,
        )

        if move_done:
            move_to_output(file_path, output_folder)

    except Exception as e:
        print(f"Error processing file: {file_path}")
        print(f"Reason: {e}")


def main():
    parser = argparse.ArgumentParser(
        description="YaPro GC335/GC337 Tool"
    )

    parser.add_argument(
        "-i",
        "--input",
        help="Input Excel file path"
    )

    parser.add_argument(
        "--folder",
        help="Process newest xlsx file inside folder"
    )

    parser.add_argument(
        "--show-browser",
        action="store_true",
        help="Show Chrome browser"
    )

    parser.add_argument(
        "--no-move-done",
        action="store_true",
        help="Do not move completed files"
    )

    args = parser.parse_args()

    if args.folder:
        run_folder_mode(
            folder_path=args.folder,
            headless=not args.show_browser,
            move_done=not args.no_move_done,
        )

    else:
        if not args.input:
            raise ValueError(
                "Missing input file. Use -i or --folder."
            )

        sheet_name = get_default_sheet_name(args.input)

        run_job(
            excel_path=args.input,
            sheet_name=sheet_name,
            headless=not args.show_browser,
            progress_cb=None,
        )


if __name__ == "__main__":
    main()