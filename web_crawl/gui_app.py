# gui_app.py

import os
import sys
import threading
import traceback
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook

import main

APP_VERSION = "1.3.0"


def resource_path(filename: str) -> str:
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        base_dir = sys._MEIPASS
    elif getattr(sys, "frozen", False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))

    return os.path.join(base_dir, filename)


def list_sheets(xlsx_path: str):
    wb = load_workbook(xlsx_path, read_only=True)
    return wb.sheetnames


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Yapro GC335/GC337 Tool")
        self.geometry("820x430")

        self.plates_file = tk.StringVar()
        self.plates_sheet = tk.StringVar()

        self.progress_var = tk.DoubleVar(value=0)
        self.progress_text = tk.StringVar(value="0 / 0")

        self.notice_img = None

        self._build_ui()

    def _build_ui(self):
        pad = {"padx": 8, "pady": 6}

        frm = ttk.Frame(self)
        frm.pack(fill="both", expand=True, **pad)

        caution_frame = ttk.LabelFrame(frm, text="⚠ Important Notice")
        caution_frame.grid(row=0, column=0, columnspan=4, sticky="we", padx=8, pady=6)

        caution_inner = ttk.Frame(caution_frame)
        caution_inner.grid(row=0, column=0, sticky="we", padx=8, pady=6)
        caution_inner.columnconfigure(0, weight=1)

        caution_text = (
            "• 執行前確保選取檔案已關閉\n"
            "• 程式會自動建立原始檔案副本.\n"
            "• 執行前請先閱讀操作說明.\n"
            "• 請確定車牌欄位在 Column D.\n"
            "• 若須查詢大量車牌，請允許一個小時以上的執行時間."
        )

        ttk.Label(
            caution_inner,
            text=caution_text,
            foreground="red",
            justify="left"
        ).grid(row=0, column=0, sticky="w")

        try:
            logo_path = resource_path("yapro_logo.png")
            self.notice_img = tk.PhotoImage(file=logo_path)
            ttk.Label(caution_inner, image=self.notice_img).grid(
                row=0, column=1, sticky="e", padx=(12, 0)
            )
        except Exception:
            pass

        caution_frame.columnconfigure(0, weight=1)

        ttk.Label(frm, text="檔案:").grid(row=1, column=0, sticky="w", **pad)
        ttk.Entry(frm, textvariable=self.plates_file, width=70).grid(row=1, column=1, sticky="we", **pad)
        ttk.Button(frm, text="選取檔案...", command=self.pick_plates_file).grid(row=1, column=2, **pad)

        ttk.Label(frm, text="工作表:").grid(row=2, column=0, sticky="w", **pad)
        self.cbo_plates_sheet = ttk.Combobox(frm, textvariable=self.plates_sheet, state="readonly", width=30)
        self.cbo_plates_sheet.grid(row=2, column=1, sticky="w", **pad)

        self.btn_run = ttk.Button(frm, text="執行", command=self.on_run)
        self.btn_run.grid(row=3, column=0, sticky="w", **pad)

        self.lbl_status = ttk.Label(frm, text="Ready.")
        self.lbl_status.grid(row=3, column=1, sticky="w", **pad)

        self.pbar = ttk.Progressbar(
            frm,
            orient="horizontal",
            mode="determinate",
            variable=self.progress_var,
            maximum=100,
        )
        self.pbar.grid(row=3, column=2, sticky="we", **pad)

        self.lbl_progress = ttk.Label(frm, textvariable=self.progress_text, width=12, anchor="e")
        self.lbl_progress.grid(row=3, column=3, sticky="e", **pad)

        self.txt = tk.Text(frm, height=14)
        self.txt.grid(row=4, column=0, columnspan=4, sticky="nsew", **pad)

        self.lbl_version = ttk.Label(frm, text=f"Version {APP_VERSION}", foreground="gray")
        self.lbl_version.grid(row=5, column=3, sticky="e", padx=8, pady=4)

        frm.columnconfigure(1, weight=1)
        frm.columnconfigure(2, weight=1)
        frm.rowconfigure(4, weight=1)

    def log(self, s: str):
        self.txt.insert("end", s + "\n")
        self.txt.see("end")

    def pick_plates_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not path:
            return

        self.plates_file.set(path)

        try:
            sheets = list_sheets(path)
            self.cbo_plates_sheet["values"] = sheets
            if sheets:
                self.plates_sheet.set(sheets[0])
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read sheets:\n{e}")

    def on_run(self):
        if not self.plates_file.get().strip():
            messagebox.showwarning("Missing", "Please choose Excel file.")
            return

        if not self.plates_sheet.get().strip():
            messagebox.showwarning("Missing", "Please choose sheet.")
            return

        self.progress_var.set(0)
        self.progress_text.set("0 / 0")

        self.btn_run.config(state="disabled")
        self.lbl_status.config(text="執行中...")
        self.log("=== START ===")

        t = threading.Thread(target=self._run_worker, daemon=True)
        t.start()

    def _run_worker(self):
        try:
            self.log(
                f"File: {self.plates_file.get()} | "
                f"sheet={self.plates_sheet.get()} | "
                f"col=D"
            )
            self.log("Output: Same file")
            self.log("Backup: Auto create _copy file")
            self.log("Mode: Auto detect from first plate")

            def progress_cb(current: int, total: int, plate: str):
                def _ui():
                    if total <= 0:
                        self.progress_var.set(0)
                        self.progress_text.set("0 / 0")
                        return

                    if current <= 0:
                        self.progress_var.set(0)
                        self.progress_text.set(f"0 / {total}")
                        return

                    pct = (current / total) * 100.0
                    self.progress_var.set(pct)
                    self.progress_text.set(f"{current} / {total}")

                    if plate:
                        self.lbl_status.config(text=f"Running... ({plate})")

                self.after(0, _ui)

            main.run_job(
                excel_path=self.plates_file.get(),
                sheet_name=self.plates_sheet.get(),
                headless=True,
                progress_cb=progress_cb,
            )

            self.log("=== DONE (saved to same file) ===")
            self._ui_done("Done.")

        except Exception:
            err = traceback.format_exc()
            self.log(err)
            self._ui_done("Error (see log).")

    def _ui_done(self, status_text: str):
        def _():
            self.lbl_status.config(text=status_text)
            self.btn_run.config(state="normal")

        self.after(0, _)


if __name__ == "__main__":
    App().mainloop()