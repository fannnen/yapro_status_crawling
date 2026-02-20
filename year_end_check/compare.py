from openpyxl import load_workbook, Workbook
from datetime import datetime, date

TOTAL_FILE = "total_sheet.xlsx"
TOTAL_SHEET = "Sheet1"
COMPARE_COL = "J"
SKIP_COL = "S"
DATE_COL = "I"

YAPRO_FILE = "yapro_sheet.xlsx"
YAPRO_SHEET = "小車"
YAPRO_COL = "E"

OUT_FILE = "unsolved_sheet.xlsx"
CUTOFF_DATE = date(2025, 9, 7)


def normalize(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def should_skip_prefix(value):
    if not value:
        return False
    v = value.upper()
    return v.startswith("E") or v.startswith("RE")


def parse_excel_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        try:
            return datetime.strptime(v.strip(), "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def read_yapro_set():
    wb = load_workbook(YAPRO_FILE, data_only=True)
    ws = wb[YAPRO_SHEET]

    result = set()
    for r in range(1, ws.max_row + 1):
        v = normalize(ws[f"{YAPRO_COL}{r}"].value)
        if v:
            result.add(v)

    return result


def main():
    total_wb = load_workbook(TOTAL_FILE, data_only=True)
    total_ws = total_wb[TOTAL_SHEET]

    yapro_set = read_yapro_set()

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Sheet1"

    # ---- copy header row first ----
    max_col = total_ws.max_column
    for c in range(1, max_col + 1):
        out_ws.cell(row=1, column=c, value=total_ws.cell(row=1, column=c).value)

    out_row = 2
    seen = set()

    # ---- process data rows ----
    for r in range(2, total_ws.max_row + 1):
        # rule 1: prefix filter
        skip_val = normalize(total_ws[f"{SKIP_COL}{r}"].value)
        if should_skip_prefix(skip_val):
            continue

        # rule 2: date filter
        date_val = parse_excel_date(total_ws[f"{DATE_COL}{r}"].value)
        if date_val and date_val < CUTOFF_DATE:
            continue

        # compare value
        compare_val = normalize(total_ws[f"{COMPARE_COL}{r}"].value)
        if not compare_val:
            continue

        if compare_val in yapro_set or compare_val in seen:
            continue

        seen.add(compare_val)

        # ---- copy entire row ----
        for c in range(1, max_col + 1):
            out_ws.cell(
                row=out_row,
                column=c,
                value=total_ws.cell(row=r, column=c).value
            )

        out_row += 1

    out_wb.save(OUT_FILE)
    print(f"✅ Done! {out_row - 2} full rows written to {OUT_FILE}")


if __name__ == "__main__":
    main()
